import os
import tempfile
import unittest

import openpyxl
from openpyxl.styles import Border, Side

import render_libro

MED, FINO = Side(style="medium"), Side(style="thin")


def _estilo(lado):
    return lado.style if lado is not None else None


def _molde(ruta):
    """Molde con los bordes irregulares del libro real: la fila 2 del martes tiene líneas finas o abiertas."""
    wb = openpyxl.Workbook()
    wb.active.title = "otra"
    for nombre in render_libro.HOJAS.values():
        ws = wb.create_sheet(nombre)
        ws.cell(60, 25).value = "x"  # fuerza el tamaño de la hoja
        for c0 in render_libro.PRIMERA_COLUMNA:
            for k in range(3):
                ws.cell(2, c0 + k).border = Border(left=MED, right=MED, top=MED, bottom=MED)
        ws.cell(2, 6).border = Border(left=FINO, right=FINO, top=FINO)      # F2: fino y sin borde inferior
        ws.cell(2, 5).border = Border(left=MED, top=MED)                    # E2: sin lado derecho ni inferior
    wb.save(ruta)


class BordesUniformes(unittest.TestCase):
    def test_todas_las_celdas_usadas_se_cierran_igual(self):
        with tempfile.TemporaryDirectory() as d:
            molde, salida = os.path.join(d, "m.xlsx"), os.path.join(d, "s.xlsx")
            _molde(molde)
            listas = {t: [[("PEREZ, Ana", "Profesora")] * 3 for _ in range(5)] for t in render_libro.HOJAS}
            render_libro.render(molde, salida, listas)
            wb = openpyxl.load_workbook(salida)
            for nombre in render_libro.HOJAS.values():
                ws = wb[nombre]
                for c0 in render_libro.PRIMERA_COLUMNA:
                    for fila in range(render_libro.PRIMERA_FILA, render_libro.PRIMERA_FILA + render_libro.RENGLONES_PRIMERA_PAGINA):
                        for k in range(3):
                            b = ws.cell(fila, c0 + k).border
                            lados = [_estilo(b.left), _estilo(b.right), _estilo(b.top), _estilo(b.bottom)]
                            self.assertEqual(lados, ["medium"] * 4, f"{nombre} fila {fila} columna {c0 + k}")
                        self.assertIsNone(_estilo(ws.cell(fila, c0 + 3).border.left))  # la columna de firma queda sin bordes

    def test_el_ultimo_renglon_de_un_dia_largo_tambien_cierra(self):
        with tempfile.TemporaryDirectory() as d:
            molde, salida = os.path.join(d, "m.xlsx"), os.path.join(d, "s.xlsx")
            _molde(molde)
            listas = {t: [[("PEREZ, Ana", "Profesora")] * 30 for _ in range(5)] for t in render_libro.HOJAS}
            render_libro.render(molde, salida, listas)
            ws = openpyxl.load_workbook(salida)["prof. tm 2026"]
            ultima = render_libro.PRIMERA_FILA + 29
            for c0 in render_libro.PRIMERA_COLUMNA:
                for k in range(3):
                    self.assertEqual(_estilo(ws.cell(ultima, c0 + k).border.bottom), "medium")
                    self.assertIsNone(_estilo(ws.cell(ultima + 1, c0 + k).border.bottom))


class Rotulos(unittest.TestCase):
    def test_rotulos_de_dia_parejos(self):
        with tempfile.TemporaryDirectory() as d:
            molde, salida = os.path.join(d, "m.xlsx"), os.path.join(d, "s.xlsx")
            _molde(molde)
            wb = openpyxl.load_workbook(molde)
            for nombre in render_libro.HOJAS.values():
                wb[nombre]["A1"], wb[nombre]["E1"] = "lunes", "    Martes"
            wb.save(molde)
            listas = {t: [[] for _ in range(5)] for t in render_libro.HOJAS}
            render_libro.render(molde, salida, listas)
            wb = openpyxl.load_workbook(salida)
            for nombre in render_libro.HOJAS.values():
                celdas = [wb[nombre].cell(1, c0) for c0 in render_libro.PRIMERA_COLUMNA]
                self.assertEqual([c.value for c in celdas], ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"])
                self.assertEqual({(c.font.sz, c.font.b, c.alignment.horizontal, c.alignment.vertical) for c in celdas},
                                 {(12.0, True, "left", "center")})


class AreaDeImpresion(unittest.TestCase):
    def test_tm_no_imprime_la_columna_en_blanco_despues_del_viernes(self):
        with tempfile.TemporaryDirectory() as d:
            molde, salida = os.path.join(d, "m.xlsx"), os.path.join(d, "s.xlsx")
            _molde(molde)
            listas = {t: [[("PEREZ, Ana", "Profesora")] * 3 for _ in range(5)] for t in render_libro.HOJAS}
            render_libro.render(molde, salida, listas)
            wb = openpyxl.load_workbook(salida)
            self.assertTrue(wb["prof. tm 2026"].print_area.endswith("$V$60"))  # el molde de prueba llega a la fila 60
            self.assertFalse(wb["prof.tmyt.2026"].print_area)


if __name__ == "__main__":
    unittest.main()
